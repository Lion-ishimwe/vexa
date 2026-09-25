"""
Add the rows of the "New products" sheet to the website catalog.

    python scripts/import_new_products.py "D:/Projects/Vexa/VEXA Products - Prices & Stock.xlsx"
    python scripts/import_new_products.py <xlsx> --dry-run
    python scripts/import_new_products.py <xlsx> --update-prices [--hold 52,53]

Reads rows from row 6 down (row 5 is the example), skips products already in
src/_data/catalog.json (same title), and appends the rest at the end so every
existing SKU stays the same. Stock quantities are never copied to the site.
Prints the new SKU for each product added.

--update-prices also copies the price of rows already on the site.
--hold lists sheet rows whose price should not be copied (e.g. a suspected typo).
"""
import json
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
CATALOG = os.path.join(ROOT, "src", "_data", "catalog.json")
FIRST_ROW = 6

# Brand spelling as it should appear on the site
BRANDS = {
    "LOGITECH": "Logitech", "EPSON": "Epson", "LENOVO": "Lenovo", "CISCO": "Cisco", "JABRA": "Jabra",
    "LIGHT WAVE": "Light Wave", "POLY": "Poly", "ASUS": "ASUS", "HP": "HP", "MSI": "MSI",
}


def slugify(s):
    s = s.lower().replace('"', "").replace("°", "").replace("&", " and ")
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")


def clean_text(s):
    """Site style: no dashes used as punctuation, brand names in normal case."""
    s = re.sub(r"\s+[-–—]\s+", ", ", str(s).strip())
    s = re.sub(r"\s+", " ", s)
    for upper, proper in BRANDS.items():
        s = re.sub(rf"^{re.escape(upper)}\b", proper, s)
    return s


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    xlsx, dry = sys.argv[1], "--dry-run" in sys.argv
    update_prices = "--update-prices" in sys.argv
    hold = set()
    if "--hold" in sys.argv:
        hold = {int(x) for x in sys.argv[sys.argv.index("--hold") + 1].split(",") if x.strip()}
    ws = load_workbook(xlsx, data_only=True)["New products"]
    cat = json.load(open(CATALOG, encoding="utf8"))
    dept_by_name = {c["name"].lower(): c for c in cat["categories"]}
    have_titles = {p["title"].lower() for p in cat["products"]}
    by_title = {p["title"].lower(): p for p in cat["products"]}
    repriced = []
    have_slugs = {p["slug"] for p in cat["products"]}

    added, skipped = [], []
    for r in range(FIRST_ROW, ws.max_row + 1):
        dept, aisle, brand, model, title, desc, price = (ws.cell(r, c).value for c in range(1, 8))
        if not title:
            continue
        dept_c = dept_by_name.get(str(dept or "").strip().lower())
        if not dept_c:
            skipped.append((r, title, f"unknown department '{dept}'"))
            continue
        title = clean_text(title)
        if title.lower() in have_titles:
            existing = by_title.get(title.lower())
            if update_prices and existing:
                try:
                    new_price = int(round(float(price))) if price not in (None, "") and float(price) > 0 else None
                except (TypeError, ValueError):
                    new_price = existing["price"]
                if new_price != existing["price"]:
                    if r in hold:
                        skipped.append((r, title, f"price held at {existing['price']:,} (sheet says {new_price:,})"))
                    else:
                        repriced.append((r, title, existing["price"], new_price))
                        existing["price"] = new_price
                continue
            skipped.append((r, title, "already on the site"))
            continue
        brand = BRANDS.get(str(brand or "").strip().upper(), str(brand or "").strip())
        slug = slugify(title)
        while slug in have_slugs:
            slug += "-2"
        try:
            amount = int(round(float(price))) if price not in (None, "") and float(price) > 0 else None
        except (TypeError, ValueError):
            amount = None
        product = {
            "slug": slug, "title": title, "brand": brand, "model": str(model or "").strip(),
            "category": dept_c["slug"], "aisle": clean_text(aisle or "Other"),
            "price": amount, "desc": clean_text(desc or ""), "img": "",
        }
        cat["products"].append(product)
        have_titles.add(title.lower())
        have_slugs.add(slug)
        added.append((r, product, "VX-" + str(len(cat["products"]) + 1000)))

    # department counts and sections follow the products
    for c in cat["categories"]:
        items = [p for p in cat["products"] if p["category"] == c["slug"]]
        c["count"] = len(items)
        c["aisles"] = list(dict.fromkeys(p["aisle"] for p in items))

    for r, p, sku in added:
        print(f"row {r:3}  {sku}  {p['category']:12} {p['aisle']:22} {p['title']}  ({p['price'] or 'price on request'})")
    for r, t, old, new in repriced:
        print(f"row {r:3}  PRICE  {t}: {old or 0:,} -> {new or 0:,}")
    for r, t, why in skipped:
        print(f"row {r:3}  SKIPPED  {t}: {why}")
    print(f"{len(added)} to add, {len(repriced)} repriced, {len(skipped)} skipped" + (" (dry run, nothing written)" if dry else ""))
    if not dry and (added or repriced):
        with open(CATALOG, "w", encoding="utf8") as f:
            f.write(json.dumps(cat, indent=2, ensure_ascii=False) + "\n")


if __name__ == "__main__":
    main()
